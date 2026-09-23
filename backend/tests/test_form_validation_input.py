"""Verificaciones de entrada y reporte del módulo Form Validation."""

import io

from openpyxl import Workbook, load_workbook

from backend.app.modules.form_validation.country import infer_country, normalize_allowed_url
from backend.app.modules.form_validation.report_service import FormValidationReportService
from backend.app.modules.weekly_auto.weekly_leads import WeeklyLeadsSpreadsheetService


def _xlsx(headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QA"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_country_inference_follows_url_rules():
    assert infer_country("", "https://utel.edu.mx/colombia/carrera") == "Colombia"
    assert infer_country("Global", "https://utel.edu.mx/argentina/maestria") == "Argentina"
    assert infer_country(None, "https://utlenlinea.com/landing") == "Perú"
    assert infer_country("", "https://utel.edu.mx/") == "México"
    assert infer_country("", "https://utel.edu.mx/indonesia") == "Indonesia"
    assert infer_country("PE", "https://universidad.utlenlinea.com/landing") == "Perú"
    assert infer_country("Global", "https://utel.edu.mx/global/", "Filipinas Bachelor") == "Filipinas"
    assert infer_country("Global", "https://utel.edu.mx/global/", "India Master") == "India"


def test_form_validation_accepts_url_only_and_omits_existing_lead():
    content = _xlsx(
        ["URL", "Country", "Lead"],
        [
            ["utel.edu.mx/colombia/carrera", "", ""],
            ["https://example.org/not-utel", "", ""],
            ["https://utel.edu.mx/usa/maestria", "USA", "https://crm.test/1"],
        ],
    )
    service = WeeklyLeadsSpreadsheetService()
    preview = service.preview(content, "urls.xlsx")
    assert preview["sheets"][0]["total_rows"] == 1
    assert preview["sheets"][0]["invalid_rows"][0]["row_number"] == 3
    rows = service.rows_for_mapping(content, preview["sheets"][0]["mapping"])
    assert rows[0]["country"] == "Colombia"
    assert rows[0]["utel_url"] == "https://utel.edu.mx/colombia/carrera"


def test_form_validation_report_has_parallel_crm_columns():
    content = _xlsx(["URL"], [["https://utel.edu.mx/"]])
    row = {"sheet": "QA", "row_number": 2, "utel_url": "https://utel.edu.mx/"}
    result = {
        "status": "PASS",
        "workflow_mode": "form_validation",
        "country": "México",
        "lead_name": "QA",
        "lead_email": "qa@example.test",
        "lead_phone": "5551234567",
        "level": "Licenciatura",
        "selected_program_name": "Administración",
        "utel_submission": "success",
        "found_inconcert": True,
        "inconcert_lead_url": "https://crm.test/inconcert/1",
        "found_balancer": False,
        "balancer_lead_url": None,
        "source_final": "inconcert",
        "search_duration_seconds": 4.2,
        "stages": [],
    }
    workbook = FormValidationReportService().build(content, {"utel_url": "URL"}, [{"row": row, "result": result}])
    saved = io.BytesIO()
    workbook.save(saved)
    output = load_workbook(io.BytesIO(saved.getvalue()), data_only=True)
    headers = [cell.value for cell in output["QA"][1]]
    assert "Link inConcert" in headers
    assert "Encontrado en Balancer" in headers
    assert output["QA"].cell(2, headers.index("Fuente final") + 1).value == "inconcert"
    ordered = [cell.value for cell in output["Form Validation resultados"][1]]
    assert ordered == list(FormValidationReportService.COLUMNS)


def test_form_validation_report_marks_fill_only_without_crm_link():
    content = _xlsx(["URL"], [["https://utel.edu.mx/"]])
    row = {"sheet": "QA", "row_number": 2, "utel_url": "https://utel.edu.mx/"}
    result = {
        "status": "PASS",
        "workflow_mode": "form_validation",
        "fill_only": True,
        "country": "México",
        "lead_name": "QA",
        "lead_email": "qa@example.test",
        "lead_phone": "5551234567",
        "level": "Licenciatura",
        "selected_program_name": "Administración",
        "utel_submission": "skipped",
        "utel_submission_message": "No enviado: modo solo llenado activo.",
        "found_inconcert": False,
        "found_balancer": False,
        "stages": [],
    }
    workbook = FormValidationReportService().build(content, {"utel_url": "URL"}, [{"row": row, "result": result}])
    saved = io.BytesIO()
    workbook.save(saved)
    output = load_workbook(io.BytesIO(saved.getvalue()), data_only=True)
    headers = [cell.value for cell in output["Form Validation resultados"][1]]
    assert output["Form Validation resultados"].cell(2, headers.index("Estado formulario") + 1).value == "Rellenado sin envío"
    qa_headers = [cell.value for cell in output["QA"][1]]
    assert output["QA"].cell(2, qa_headers.index("RESULTADO FORMULARIO") + 1).value == "RELLENADO - NO ENVIADO"


def test_report_landing_columns_are_detected_and_normalized():
    content = _xlsx(
        ["pais", "url", "landing", "carga_formularios_utel", "integracion_envio", "Lead"],
        [["PE", "https://universidad.utlenlinea.com/experiencias", "universidad/experiencias", "SI (activo render/create)", "Balanceador", ""]],
    )
    service = WeeklyLeadsSpreadsheetService()
    preview = service.preview(content, "Reporte Landings.xlsx")
    sheet = preview["sheets"][0]
    assert sheet["mapping"]["country"] == "pais"
    assert sheet["mapping"]["utel_url"] == "url"
    assert sheet["mapping"]["form_signal"] == "carga_formularios_utel"
    rows = service.rows_for_mapping(content, sheet["mapping"])
    assert rows[0]["country"] == "Perú"
    assert rows[0]["weekly_form_type"] == "form_lp"
    assert rows[0]["integration_hint"] == "Balanceador"


def test_form_validation_deduplicates_report_sheets_but_retries_non_url_notes():
    """Las hojas espejo no repiten filas y los comentarios no bloquean reintentos."""

    workbook = Workbook()
    clean = workbook.active
    clean.title = "CLEAN"
    clean.append(["pais", "url", "Lead"])
    clean.append(["PE", "https://utlenlinea.com/landing-a", "https://crm.test/1"])
    clean.append(["PE", "https://utlenlinea.com/landing-b", "no funciona boton"])
    mirror = workbook.create_sheet("Hoja 1")
    mirror.append(["pais", "url"])
    mirror.append(["PE", "https://utlenlinea.com/landing-a"])
    mirror.append(["PE", "https://utlenlinea.com/landing-b"])
    output = io.BytesIO()
    workbook.save(output)

    service = WeeklyLeadsSpreadsheetService()
    content = output.getvalue()
    preview = service.preview(content, "reporte.xlsx")
    assert preview["total_rows"] == 1
    rows = service.rows_for_mapping(content, preview["sheets"][0]["mapping"])
    assert len(rows) == 1
    assert rows[0]["utel_url"] == "https://utlenlinea.com/landing-b"
