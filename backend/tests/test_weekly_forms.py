import asyncio
import io
from pathlib import Path
from unittest.mock import AsyncMock, Mock

from openpyxl import Workbook, load_workbook

from backend.app.config.settings import Settings
from backend.app.modules.weekly_auto.weekly_forms import (
    WeeklyFormsCaseConfig,
    WeeklyFormsRunner,
    WeeklyFormsSpreadsheetService,
)
from backend.app.modules.bot_leads_deploy.runner import UtelQaError
from backend.app.services.bot_report_service import BotReportService


def _workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QA"
    sheet.append(["Country", "Nivel", "Activo de Test", "Location", "Cliente", "Lead"])
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_weekly_forms_uses_displayed_url_infers_level_and_skips_existing_lead():
    content = _workbook_bytes([
        ["USA", "", "https://universidad.utel.edu.mx/usa/licenciaturas-online", "Form Lp", "LatAm", ""],
        ["Global", "Licenciatura", "https://educacioncontinua.utel.mx/landing/educacion-continua", "Form Lp", "Mex", ""],
        ["México", "Licenciatura", "https://example.test/already", "Footer", "Mex", "https://crm.test/lead/1"],
        ["México", "Licenciatura", "https://example.test/senior", "Targeta", "Mex", ""],
    ])
    service = WeeklyFormsSpreadsheetService()
    preview = service.preview(content, "QA.xlsx")
    rows = service.rows_for_mapping(content, preview["sheets"][0]["mapping"])

    assert len(rows) == 3
    assert rows[0]["level"] == "Licenciatura"
    assert rows[1]["utel_url"].endswith("/landing/educacion-continua")
    assert rows[2]["weekly_form_type"] == "form_lp"
    assert service.BATCH_SIZE == 5
    assert service.BATCH_PAUSE_SECONDS == 60


def test_weekly_forms_preserves_cell_value_instead_of_hyperlink_target():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Country", "Nivel", "Activo de Test", "Location", "Cliente", "Lead"])
    visible = "https://educacioncontinua.utel.mx/landing/educacion-continua"
    sheet.append(["Global", "Licenciatura", visible, "Form Lp", "Mex", ""])
    sheet.cell(2, 3).hyperlink = visible + ".html"
    output = io.BytesIO()
    workbook.save(output)
    service = WeeklyFormsSpreadsheetService()
    preview = service.preview(output.getvalue(), "QA.xlsx")
    rows = service.rows_for_mapping(output.getvalue(), preview["sheets"][0]["mapping"])
    assert rows[0]["utel_url"] == visible


def test_weekly_forms_requires_lead_output_column():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Country", "Nivel", "Activo de Test", "Location", "Cliente"])
    sheet.append(["USA", "Licenciatura", "https://example.test/form", "Form Lp", "LatAm"])
    output = io.BytesIO()
    workbook.save(output)

    assert WeeklyFormsSpreadsheetService().preview(output.getvalue(), "QA.xlsx")["sheets"] == []


def test_manual_form_failure_keeps_lead_cell_blank():
    content = _workbook_bytes([
        ["USA", "Licenciatura", "https://example.test/no-form", "Form Lp", "LatAm", ""],
    ])
    service = WeeklyFormsSpreadsheetService()
    preview = service.preview(content, "QA.xlsx")
    mapping = preview["sheets"][0]["mapping"]
    row = service.rows_for_mapping(content, mapping)[0]
    result = {
        "status": "FAIL",
        "dry_run": False,
        "lead_url": None,
        "lead_email": "persona@example.test",
        "selected_program_name": "",
        "utel_submission_attempted": False,
        "stages": [{"stage": "weekly_manual", "status": "FAIL", "message": "Sin formulario"}],
        "summary": "Completar manualmente",
    }

    workbook = BotReportService().build(content, mapping, [{"row": row, "result": result}])
    saved = io.BytesIO()
    workbook.save(saved)
    output = load_workbook(io.BytesIO(saved.getvalue()), data_only=True)
    assert output["QA"].cell(2, 6).value is None


def test_generic_lp_form_is_identified_and_filled(tmp_path: Path):
    async def scenario():
        from playwright.async_api import async_playwright

        settings = Settings(database_path=tmp_path / "test.db", storage_dir=tmp_path / "storage")
        runner = WeeklyFormsRunner(settings)
        config = WeeklyFormsCaseConfig.model_validate({
            "name": "Weekly Forms test",
            "environment": "sandbox",
            "dry_run": True,
            "country": "USA",
            "utel_url": "https://example.test/form",
            "modality": "En linea",
            "level": "Licenciatura",
            "form_type": "lateral",
            "weekly_form_type": "form_lp",
            "workflow_mode": "form_validation",
            "lead": {"name": "Persona QA", "email": "persona@example.test", "phone": "2125550199"},
        })
        runner._rotation_config = config
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(headless=True)
            page = await browser.new_page()
            await page.set_content("""
              <form id="lead-form">
                <label>Nombre <input name="name" required></label>
                <label>Email <input name="email" type="email" required></label>
                <label>Teléfono <input name="phone" type="tel" required></label>
                <label>Área de interés <select name="area" required><option value="">Selecciona</option><option value="lic">Licenciatura</option></select></label>
                <label>Programa <select name="program" required><option value="">Selecciona</option><option value="adm">Administración</option></select></label>
                <label>Privacidad <input name="privacy" type="checkbox" required></label>
                <button type="submit">Solicitar información</button>
              </form>
            """)
            form = await runner._find_utel_form(page, config)
            await runner._fill_utel_form(page, form, config)
            assert await form.locator('[name="name"]').input_value() == "Persona QA"
            assert await form.locator('[name="email"]').input_value() == "persona@example.test"
            assert await form.locator('[name="phone"]').input_value() == "2125550199"
            assert await form.locator('[name="area"]').input_value() == "lic"
            assert await form.locator('[name="program"]').input_value() == "adm"
            assert await form.locator('[name="privacy"]').is_checked()
            await browser.close()

    asyncio.run(scenario())


def test_page_without_lead_form_is_reported_for_manual_work(tmp_path: Path):
    async def scenario():
        from playwright.async_api import async_playwright

        runner = WeeklyFormsRunner(Settings(database_path=tmp_path / "test.db", storage_dir=tmp_path / "storage"))
        runner.GENERIC_FORM_TIMEOUT_MS = 50
        config = WeeklyFormsCaseConfig.model_validate({
            "country": "Mexico", "utel_url": "https://example.test/no-form",
            "modality": "En linea", "level": "Licenciatura", "form_type": "lateral",
            "weekly_form_type": "form_lp", "workflow_mode": "form_validation",
            "lead": {"name": "Persona QA", "email": "persona@example.test", "phone": "5551234567"},
        })
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(headless=True)
            page = await browser.new_page()
            await page.set_content("<main><h1>Contenido sin formulario</h1></main>")
            try:
                await runner._find_utel_form(page, config)
                raise AssertionError("Se esperaba UtelQaError")
            except UtelQaError as error:
                assert error.stage == "weekly_manual"
                assert "se dejará en blanco" in str(error)
            await browser.close()

    asyncio.run(scenario())


def test_custom_privacy_checkbox_is_activated(tmp_path: Path):
    async def scenario():
        from playwright.async_api import async_playwright

        runner = WeeklyFormsRunner(Settings(database_path=tmp_path / "test.db", storage_dir=tmp_path / "storage"))
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(headless=True)
            page = await browser.new_page()
            await page.set_content("""
              <input id="terms" type="checkbox">
              <label for="terms">Acepto la política de privacidad</label>
              <script>document.querySelector('#terms').addEventListener('click', event => event.preventDefault())</script>
            """)
            checkbox = page.locator("#terms")
            await runner._ensure_checkbox_checked(checkbox)
            assert await checkbox.is_checked()
            await browser.close()

    asyncio.run(scenario())


def test_parallel_crm_uses_the_first_detail_as_primary_result(tmp_path: Path):
    """Weekly Forms conserva como enlace principal el CRM que responde primero."""

    async def scenario():
        runner = WeeklyFormsRunner(Settings(database_path=tmp_path / "test.db", storage_dir=tmp_path / "storage"))
        runner._submission_attempted = True
        inconcert_url = "https://crm.test/inconcert/contact/1"
        balancer_url = "https://lead-balancer.scalahed.com/leads/detail/2"

        # Las páginas aisladas representan las dos pestañas que usa el flujo
        # paralelo. El detalle del Balanceador termina antes que InConcert.
        inconcert_page = Mock(url="https://crm.test/inconcert/contacts")
        balancer_page = Mock(url="https://lead-balancer.scalahed.com/leads/")
        inconcert_page.set_default_timeout = Mock()
        balancer_page.set_default_timeout = Mock()
        context = Mock(new_page=AsyncMock(side_effect=[inconcert_page, balancer_page]))

        async def run_stage(_number, _stage, _message, _page, action, _screenshot=None):
            return await action()

        async def search_inconcert(_page, _email, _name):
            await asyncio.sleep(0.04)

        async def search_balancer(_page, _email, _name):
            await asyncio.sleep(0.01)
            runner.lead_url = balancer_url

        async def open_manage(_page, _name, _email):
            await asyncio.sleep(0.02)
            runner.lead_url = inconcert_url
            return Mock(url=inconcert_url)

        runner._run_stage = run_stage
        runner._open_inconcert = AsyncMock()
        runner._login_inconcert = AsyncMock()
        runner._open_contacts = AsyncMock()
        runner._search_lead = search_inconcert
        runner._search_lead_balancer = search_balancer
        runner._open_manage = open_manage

        config = WeeklyFormsCaseConfig.model_validate({
            "name": "Weekly Forms first CRM",
            "environment": "sandbox",
            "country": "Mexico",
            "utel_url": "https://utel.edu.mx/programa",
            "inconcert_url": "https://crm.test",
            "modality": "En linea",
            "level": "Licenciatura",
            "workflow_mode": "form_validation",
            "lead": {"name": "Persona QA", "email": "persona@example.test", "phone": "5551234567"},
        })

        result = await runner._verify_crm_parallel(
            context,
            config,
            None,
            "2026-09-15T10:00:00",
            0.0,
        )

        assert result["inconcert_lead_url"] == inconcert_url
        assert result["balancer_lead_url"] == balancer_url
        assert result["lead_url"] == balancer_url
        assert result["lead_source"] == "balanceador"
        assert result["source_final"] == "balanceador"
        assert result["first_crm_source"] == "balanceador"

    asyncio.run(scenario())
