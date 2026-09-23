from io import BytesIO

from openpyxl import load_workbook

from backend.app.modules.strapi_products.report import build_combined_report


def test_combined_report_includes_both_processes_and_pdp_plan():
    canonical = {
        "status": "SUCCESS",
        "summary": {"total": 1, "updated": 1},
        "results": [{
            "sheet": "Productos", "row_number": 2, "program": "Licenciatura X", "country": "Argentina",
            "status": "UPDATED", "old_canonical": None, "new_canonical": "https://example.test/x",
        }],
    }
    pdp = {
        "status": "WARNING",
        "summary": {"total": 1, "failed": 1},
        "results": [{
            "sheet": "Productos", "row_number": 2, "program": "Licenciatura X", "country": "Argentina",
            "status": "FAILED", "siu_key": "ABC", "banner_key": "ABC", "description": "Descripción",
            "verification": {"ok": False}, "message": "No coincide", "changes": [
                {"field": "content", "action": "update", "before": "old", "after": "new", "verified": False},
            ],
        }],
    }

    workbook = load_workbook(BytesIO(build_combined_report(canonical, pdp)), data_only=True)

    results = workbook["Resultados"]
    assert results.max_row == 3
    assert results["A2"].value == "Canonical"
    assert results["A3"].value == "PDP"
    assert results["I3"].value == "ABC"
    assert workbook["Resumen"].max_row >= 5
    plan = workbook["Plan por producto"]
    assert plan.max_row == 2
    assert plan["A2"].value == "PDP"
    assert plan["F2"].value == "content"
