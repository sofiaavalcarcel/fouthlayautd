import asyncio
import io

from docx import Document

from backend.app.modules.strapi_products.bullet_tabs import (
    BulletTabContent,
    build_bullet_tab_payload,
    extract_bullet_tabs,
)
from backend.app.modules.strapi_products.description_runner import StrapiDescriptionRunner
from backend.app.modules.strapi_products.models import ProductRow
from backend.app.modules.strapi_products.pdp_description import (
    extract_benefits,
    extract_graduate_testimonies,
    extract_labor_field_content,
    extract_pdp_market_description,
)


def _document() -> bytes:
    document = Document()
    document.add_paragraph("Programa de prueba", style="Title")
    document.add_paragraph("Descripción general del PDP.")
    document.add_paragraph("Asignaturas", style="Heading 1")
    document.add_paragraph("Información del plan.")
    document.add_paragraph("1° cuatrimestre", style="Heading 2")
    for heading, intro, bullets in (
        ("Perfil de ingreso:", "Texto introductorio de ingreso.", ["Interés ambiental: comprender el entorno.", "Análisis: organizar información."]),
        ("Perfil de egreso", "Texto introductorio de egreso.", ["Diseño: crear proyectos."]),
        ("¿Dónde podrás trabajar?", "Texto introductorio de empleabilidad.", ["Consultoría: asesorar organizaciones."]),
    ):
        document.add_paragraph(heading, style="Heading 1")
        document.add_paragraph(intro)
        for text in bullets:
            document.add_paragraph(text, style="List Bullet")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_extract_bullet_tabs_separates_intro_from_bulleted_content():
    tabs = extract_bullet_tabs("pdp.docx", _document())
    assert tabs["Perfil ingreso"].description == "Texto introductorio de ingreso."
    assert tabs["Perfil ingreso"].bullets == (
        "Interés ambiental: comprender el entorno.",
        "Análisis: organizar información.",
    )
    assert tabs["Perfil egreso"].description == "Texto introductorio de egreso."
    assert tabs["Empleabilidad"].bullets == ("Consultoría: asesorar organizaciones.",)


def test_extract_bullet_tabs_uses_profile_table_from_pdp():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    for cell, title, bullets in (
        (table.cell(0, 0), "Perfil de ingreso", ["Interés por el programa.", "Creatividad para aprender."]),
        (table.cell(0, 1), "Perfil de egreso", ["Gestionarás proyectos.", "Medirás resultados."]),
    ):
        cell.paragraphs[0].text = title
        for bullet in bullets:
            cell.add_paragraph(bullet)
    document.add_paragraph("Empleabilidad")
    document.add_paragraph("Texto de empleabilidad.")
    document.add_paragraph("Consultoría: asesorar organizaciones.", style="List Bullet")
    output = io.BytesIO()
    document.save(output)

    tabs = extract_bullet_tabs("pdp.docx", output.getvalue())
    assert tabs["Perfil ingreso"].description == ""
    assert tabs["Perfil egreso"].description == ""
    assert tabs["Perfil ingreso"].bullets == ("Interés por el programa.", "Creatividad para aprender.")
    assert tabs["Perfil egreso"].bullets == ("Gestionarás proyectos.", "Medirás resultados.")
    assert tabs["Empleabilidad"].bullets == ("Consultoría: asesorar organizaciones.",)


def test_build_profile_tab_without_description_removes_bullets_description():
    existing = {
        "id": 25,
        "attributes": {
            "strapiName": "Perfil ingreso Programa A",
            "content": [{
                "id": 12,
                "__component": "section.bullets",
                "bulletsDescription": {"desktop": "No debe conservarse"},
                "bullets": [],
            }],
        },
    }
    payload = build_bullet_tab_payload(
        "Programa A",
        BulletTabContent("Perfil ingreso", "", ("Una viñeta del PDP.",)),
        existing=existing,
        locale="es-MX",
    )
    assert payload["content"][0]["bulletsDescription"] is None


def test_extract_benefits_from_pdp_block():
    document = Document()
    document.add_paragraph("(Padrón) Bloque. Beneficios de estudiar en Utel")
    document.add_paragraph("Título con validez oficial SEP")
    document.add_paragraph("Preparación para el mundo laboral")
    document.add_paragraph("Titulación directa")
    document.add_paragraph("(Web) Bloque. Qué es el programa")
    output = io.BytesIO()
    document.save(output)
    assert extract_benefits("pdp.docx", output.getvalue()) == (
        "Título con validez oficial SEP",
        "Preparación para el mundo laboral",
        "Titulación directa",
    )


def test_extract_graduate_testimonies_from_coms_pdp_block_with_labeled_cards():
    document = Document()
    document.add_paragraph("(Coms) Bloque. Egresados de la Licenciatura en Psicología Organizacional")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Me ayudó a crecer profesionalmente.\nAna López\nCarrera:\nEgresado"
    output = io.BytesIO()
    document.save(output)

    assert extract_graduate_testimonies(
        "pdp.docx", output.getvalue(), "Licenciatura en Psicología Organizacional"
    ) == [("Ana López", "Me ayudó a crecer profesionalmente.")]


def test_extract_graduate_testimonies_from_explicit_name_comment_labels():
    document = Document()
    document.add_paragraph("(Coms) Bloque. Egresados de la Licenciatura en Psicología Organizacional")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Nombre: Ana López\nComentario: La modalidad me permitió estudiar y trabajar."
    output = io.BytesIO()
    document.save(output)

    assert extract_graduate_testimonies(
        "pdp.docx", output.getvalue(), "Licenciatura en Psicología Organizacional"
    ) == [("Ana López", "La modalidad me permitió estudiar y trabajar.")]


def test_extract_graduate_testimonies_from_pdp_cards_with_program_and_year():
    document = Document()
    document.add_paragraph("(Coms) Bloque. Egresados de la Licenciatura en Psicología Organizacional")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = (
        "Esta carrera me ayudó a crecer profesionalmente.\n"
        "Ana López\n"
        "Licenciatura en Psicología Organizacional\n"
        "Egresada, 2024"
    )
    table.cell(0, 1).text = (
        "La carrera me dio herramientas para mejorar mi trabajo.\n"
        "Carlos Pérez\n"
        "Licenciatura en Psicología Organizacional\n"
        "Egresado, 2023"
    )
    output = io.BytesIO()
    document.save(output)

    assert extract_graduate_testimonies(
        "pdp.docx", output.getvalue(), "Licenciatura en Psicología Organizacional"
    ) == [
        ("Ana López", "Esta carrera me ayudó a crecer profesionalmente."),
        ("Carlos Pérez", "La carrera me dio herramientas para mejorar mi trabajo."),
    ]


def test_extract_active_graduates_percentage_from_banner_table():
    document = Document()
    document.add_paragraph("(Padron) Banner.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "90%\nEgresados activos en el mundo laboral"
    output = io.BytesIO()
    document.save(output)

    from backend.app.modules.strapi_products.pdp_description import extract_active_graduates_percentage

    assert extract_active_graduates_percentage("pdp.docx", output.getvalue()) == 90


def test_extract_labor_field_sections_without_mixing_areas_positions_and_market():
    document = Document()
    document.add_paragraph("(Web) Bloque. Oportunidades Profesionales")
    document.add_paragraph("Áreas en las que vas a poder trabajar")
    document.add_paragraph("Sector público: diseñar políticas y coordinar proyectos.")
    document.add_paragraph("Puestos que podrías ocupar")
    document.add_paragraph("Analista de políticas públicas")
    document.add_paragraph("Director de proyectos")
    document.add_paragraph("Datos de mercado laboral")
    document.add_paragraph("El campo laboral ofrece oportunidades en organizaciones públicas y privadas.")
    output = io.BytesIO()
    document.save(output)

    extracted = extract_labor_field_content("pdp.docx", output.getvalue())

    assert extracted.areas == (("Sector público", "diseñar políticas y coordinar proyectos."),)
    assert extracted.positions == ("Analista de políticas públicas", "Director de proyectos")
    assert extracted.market_description == "El campo laboral ofrece oportunidades en organizaciones públicas y privadas."


def test_extract_labor_field_pairs_alternating_area_titles_and_descriptions():
    document = Document()
    document.add_paragraph("Áreas en las que vas a poder trabajar")
    document.add_paragraph("Comunicación organizacional")
    document.add_paragraph("Participa en estrategias de comunicación interna y externa.")
    document.add_paragraph("Publicidad y marketing digital")
    document.add_paragraph("Desarrolla campañas y contenidos para posicionar marcas.")
    document.add_paragraph("Relaciones públicas")
    document.add_paragraph("Gestiona vínculos entre organizaciones y sus audiencias.")
    document.add_paragraph("Producción de contenidos y medios")
    document.add_paragraph("Crea proyectos escritos, audiovisuales y multimedia.")
    output = io.BytesIO()
    document.save(output)

    extracted = extract_labor_field_content("pdp.docx", output.getvalue())

    assert extracted.areas == (
        ("Comunicación organizacional", "Participa en estrategias de comunicación interna y externa."),
        ("Publicidad y marketing digital", "Desarrolla campañas y contenidos para posicionar marcas."),
        ("Relaciones públicas", "Gestiona vínculos entre organizaciones y sus audiencias."),
        ("Producción de contenidos y medios", "Crea proyectos escritos, audiovisuales y multimedia."),
    )


def test_extract_pdp_market_description_uses_only_the_coms_block():
    document = Document()
    document.add_paragraph("(Web) Datos de mercado laboral")
    document.add_paragraph("Texto de Web que no debe copiarse.")
    document.add_paragraph("(Coms) Datos de mercado laboral")
    document.add_paragraph("Texto exacto del PDP para el mercado laboral.")
    document.add_paragraph("(Web) Bloque siguiente")
    document.add_paragraph("Texto posterior que no debe copiarse.")
    output = io.BytesIO()
    document.save(output)

    assert extract_pdp_market_description("pdp.docx", output.getvalue()) == (
        "Texto exacto del PDP para el mercado laboral."
    )


def test_extract_pdp_market_description_stops_before_later_testimonials_table():
    document = Document()
    document.add_paragraph("(Coms) Datos de mercado laboral")
    document.add_paragraph("Texto exacto del mercado laboral.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = (
        "(Coms) Bloque. Egresados de Programa A\n"
        "Un testimonio que no pertenece al mercado laboral.\n"
        "Andrea Castillo\n"
        "Egresada, 2024"
    )
    output = io.BytesIO()
    document.save(output)

    assert extract_pdp_market_description("pdp.docx", output.getvalue()) == (
        "Texto exacto del mercado laboral."
    )


def test_extract_pdp_market_description_returns_empty_for_an_empty_coms_block():
    document = Document()
    document.add_paragraph("(Coms) Datos de mercado laboral")
    document.add_paragraph("(Coms) Bloque. Egresados de Programa A")
    output = io.BytesIO()
    document.save(output)

    assert extract_pdp_market_description("pdp.docx", output.getvalue()) == ""


def test_resolve_active_graduates_percentage_defaults_to_90_without_banner_value():
    document = Document()
    output = io.BytesIO()
    document.save(output)

    from backend.app.modules.strapi_products.pdp_description import resolve_active_graduates_percentage

    assert resolve_active_graduates_percentage("pdp.docx", output.getvalue()) == 90


def test_resolve_active_graduates_percentage_keeps_pdp_value_when_present():
    document = Document()
    document.add_paragraph("(Padron) Banner.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "75%\nEgresados activos en el mundo laboral"
    output = io.BytesIO()
    document.save(output)

    from backend.app.modules.strapi_products.pdp_description import resolve_active_graduates_percentage

    assert resolve_active_graduates_percentage("pdp.docx", output.getvalue()) == 75


def test_build_existing_bullet_tab_changes_only_its_own_text_and_preserves_settings():
    existing = {
        "id": 25,
        "attributes": {
            "strapiName": "Perfil ingreso Programa A",
            "title": {"id": 90, "desktop": "Perfil ingreso"},
            "content": [{
                "id": 12,
                "__component": "section.bullets",
                "iconsColor": "secondary.default",
                "hideSection": False,
                "bulletsDescription": {"id": 91, "desktop": "Texto anterior", "chakraConfig": None},
                "bullets": [{"id": 92, "title": None, "description": "Anterior", "chakraConfig": None}],
            }],
        },
    }
    template = {"id": 8, "attributes": {"strapiName": "Perfil ingreso Otro programa", "content": [{"id": 3, "__component": "section.bullets", "iconsColor": "red", "bullets": [{"id": 2, "description": "Texto prestado"}]}]}}
    payload = build_bullet_tab_payload(
        "Programa A", BulletTabContent("Perfil ingreso", "Nuevo párrafo.", ("Viñeta fiel al documento.",)),
        existing=existing, template=template, locale="es-AR",
    )
    component = payload["content"][0]
    assert component["id"] == 12
    assert component["iconsColor"] == "secondary.default"
    assert component["bulletsDescription"]["desktop"] == "Nuevo párrafo."
    assert component["bullets"][0]["description"] == "Viñeta fiel al documento."
    assert "strapiName" not in payload
    assert existing["attributes"]["content"][0]["bullets"][0]["description"] == "Anterior"
    assert template["attributes"]["content"][0]["bullets"][0]["description"] == "Texto prestado"


def test_build_new_bullet_tab_copies_style_without_reusing_component_ids():
    template = {
        "id": 8,
        "attributes": {
            "title": {"id": 77, "desktop": "Perfil ingreso", "chakraConfig": None},
            "iconTitle": None,
            "content": [{
                "id": 3,
                "__component": "section.bullets",
                "iconsColor": "secondary.default",
                "hideSection": False,
                "bulletsDescription": {"id": 4, "desktop": "Texto de referencia", "chakraConfig": None},
                "bullets": [{"id": 5, "description": "Texto ajeno", "chakraConfig": None}],
            }],
        },
    }
    payload = build_bullet_tab_payload(
        "Programa B", BulletTabContent("Perfil ingreso", "Texto del programa B.", ("Viñeta del programa B.",)),
        template=template, locale="es-AR",
    )
    assert payload["strapiName"] == "Perfil ingreso Programa B"
    assert payload["locale"] == "es-AR"
    assert payload["title"]["desktop"] == "Perfil ingreso"
    component = payload["content"][0]
    assert "id" not in component
    assert component["iconsColor"] == "secondary.default"
    assert component["bulletsDescription"]["desktop"] == "Texto del programa B."
    assert "id" not in component["bulletsDescription"]
    assert component["bullets"] == [{
        "title": None,
        "description": "Viñeta del programa B.",
        "chakraConfig": None,
        "icon": {"name": "UilCheck", "chakraConfig": None},
        "iconMobile": None,
    }]
    assert template["attributes"]["content"][0]["bullets"][0]["description"] == "Texto ajeno"


def test_existing_tab_without_desktop_image_copies_same_locale_template_media():
    existing = {
        "id": 25,
        "attributes": {
            "strapiName": "Perfil egreso Programa A",
            "content": [{
                "id": 12,
                "__component": "section.bullets",
                "coverImage": {"id": 50, "desktop": None},
                "bullets": [],
            }],
        },
    }
    template = {
        "id": 26,
        "attributes": {
            "strapiName": "Perfil egreso Programa B",
            "locale": "es-AR",
            "content": [{
                "id": 13,
                "__component": "section.bullets",
                "coverImage": {
                    "id": 51,
                    "desktop": {
                        "id": 52,
                        "image": {"data": {"id": 700, "attributes": {"name": "egreso.png"}}},
                        "objectFit": "cover",
                        "chakraConfig": {"objectPosition": "top center"},
                        "preload": False,
                    },
                },
                "bullets": [],
            }],
        },
    }
    payload = build_bullet_tab_payload(
        "Programa A", BulletTabContent("Perfil egreso", "Intro.", ("Viñeta.",)),
        existing=existing, template=template, locale="es-AR",
    )
    cover = payload["content"][0]["coverImage"]
    assert cover["id"] == 50
    assert cover["desktop"] == {
        "image": {"id": 700},
        "objectFit": "cover",
        "chakraConfig": {"objectPosition": "top center"},
        "preload": False,
    }
    assert template["attributes"]["content"][0]["coverImage"]["desktop"]["id"] == 52


def test_existing_tab_normalizes_populated_desktop_media_for_strapi_updates():
    existing = {
        "id": 25,
        "attributes": {
            "strapiName": "Perfil egreso Programa A",
            "content": [{
                "id": 12,
                "__component": "section.bullets",
                "coverImage": {
                    "id": 50,
                    "desktop": {
                        "id": 52,
                        "image": {"data": {"id": 700, "attributes": {"name": "egreso.png"}}},
                        "objectFit": "cover",
                    },
                },
                "bullets": [],
            }],
        },
    }
    payload = build_bullet_tab_payload(
        "Programa A", BulletTabContent("Perfil egreso", "Intro.", ("Viñeta.",)),
        existing=existing, locale="es-AR",
    )
    desktop = payload["content"][0]["coverImage"]["desktop"]
    assert desktop["id"] == 52
    assert desktop["image"] == {"id": 700}


def test_description_runner_updates_only_product_specific_bullet_tabs(tmp_path):
    source = tmp_path / "programa.docx"
    source.write_bytes(_document())
    updates = []
    created = []

    class FakeClient:
        def __init__(self):
            self.entries = {
                "Perfil ingreso Programa de prueba": {"id": 41, "attributes": {"strapiName": "Perfil ingreso Programa de prueba", "content": [{"id": 1, "__component": "section.bullets", "bullets": [], "bulletsDescription": {"desktop": "viejo"}}]}},
                "Perfil egreso Programa de prueba": {"id": 42, "attributes": {"strapiName": "Perfil egreso Programa de prueba", "content": [{"id": 2, "__component": "section.bullets", "bullets": [], "bulletsDescription": {"desktop": "viejo"}}]}},
                "Empleabilidad Programa de prueba": None,
            }
            self.template = {"id": 99, "attributes": {"title": {"desktop": "Empleabilidad"}, "content": [{"id": 9, "__component": "section.bullets", "iconsColor": "secondary.default", "bullets": []}]}}

        async def find_product(self, *args, **kwargs): return {"id": 7}
        async def get_product_tabs_bullet_section(self, *args):
            return {"id": 8, "hideSection": False, "tabs": [
                {"id": 41, "strapiName": "Perfil de ingreso Programa de prueba"},
                {"id": 42, "strapiName": "Perfil de egreso Programa de prueba"},
            ]}
        async def get_bullet_tab_by_id(self, identifier):
            return next(entry for entry in self.entries.values() if entry and entry["id"] == identifier)
        async def find_bullet_tab_template(self, prefix, locale): return self.template
        async def update_bullet_tab(self, identifier, payload):
            entry = next(item for item in self.entries.values() if item and item["id"] == identifier)
            entry["attributes"].update(payload)
            updates.append((identifier, payload))
            return {"id": identifier}
        async def create_bullet_tab(self, payload):
            created.append(payload)
            self.entries[payload["strapiName"]] = {"id": 43, "attributes": payload}
            return {"id": 43}
        async def update_product(self, identifier, payload): updates.append(("product", payload)); return {}

    client = FakeClient()
    results, summary = asyncio.run(StrapiDescriptionRunner(client, "Argentina", "es-AR", dry_run=False).run([
        ProductRow("Sheet", 2, "Programa de prueba", str(source))
    ]))
    assert results[0].status == "UPDATED", (results[0].message, results[0].verification)
    assert summary.updated == 1
    assert [item[0] for item in updates[:2]] == [41, 42]
    assert updates[0][1]["content"][0]["bulletsDescription"]["desktop"] == "Texto introductorio de ingreso."
    assert updates[0][1]["content"][0]["bullets"][0]["description"] == "Interés ambiental: comprender el entorno."
    assert updates[0][1]["content"][0]["bullets"][0]["icon"]["name"] == "UilCheck"
    assert len(created) == 1
    assert created[0]["strapiName"] == "Empleabilidad Programa de prueba"
    product_update = updates[-1][1]
    assert product_update["tabsBulletSection"]["tabs"] == [{"id": 41}, {"id": 42}, {"id": 43}]
    assert client.template["attributes"]["content"][0]["iconsColor"] == "secondary.default"


def test_description_runner_dry_run_does_not_create_or_update_strapi_entries(tmp_path):
    source = tmp_path / "programa.docx"
    source.write_bytes(_document())
    writes = []

    class FakeClient:
        async def find_product(self, *args, **kwargs): return {"id": 7}
        async def get_product_tabs_bullet_section(self, *args):
            return {"id": 8, "hideSection": False, "tabs": [
                {"id": 41, "strapiName": "Perfil de ingreso Programa de prueba"},
                {"id": 42, "strapiName": "Perfil de egreso Programa de prueba"},
            ]}
        async def get_bullet_tab_by_id(self, identifier):
            return {"id": identifier, "attributes": {"strapiName": f"tab {identifier}", "content": [{"id": identifier, "__component": "section.bullets", "bullets": []}]}}
        async def find_bullet_tab_template(self, prefix, locale):
            return {"id": 99, "attributes": {"title": {"desktop": prefix}, "content": [{"id": 9, "__component": "section.bullets", "bullets": []}]}}
        async def update_bullet_tab(self, *args): writes.append(("update-tab", args))
        async def create_bullet_tab(self, *args): writes.append(("create-tab", args))
        async def update_product(self, *args): writes.append(("update-product", args))

    results, summary = asyncio.run(StrapiDescriptionRunner(FakeClient(), "Argentina", "es-AR", dry_run=True).run([
        ProductRow("Sheet", 2, "Programa de prueba", str(source))
    ]))
    assert results[0].status == "DRY_RUN"
    assert summary.dry_run == 1
    assert writes == []
