import io

from docx import Document

from backend.app.modules.strapi_products.common_questions import (
    build_common_questions_payload,
    common_questions_match,
    extract_common_questions,
)
from backend.app.modules.strapi_products.description_runner import StrapiDescriptionRunner
from backend.app.modules.strapi_products.models import ProductRow
import asyncio


def _faq_doc() -> bytes:
    doc = Document()
    doc.add_heading("12. Preguntas frecuentes (FAQ)", level=2)
    doc.add_paragraph("1. ¿Cuánto dura el programa?")
    doc.add_paragraph("Dura 48 meses.")
    doc.add_paragraph("3. ¿Tiene validez oficial?")
    doc.add_paragraph("Sí, cuenta con RVOE.")
    doc.add_heading("Referencias", level=2)
    doc.add_paragraph("Texto que no pertenece a la FAQ.")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def test_extracts_numbered_faq_until_next_heading_and_keeps_exact_copy():
    entries = extract_common_questions("pdp.docx", _faq_doc())
    assert [(item.question, item.answer) for item in entries] == [
        ("¿Cuánto dura el programa?", "Dura 48 meses."),
        ("¿Tiene validez oficial?", "Sí, cuenta con RVOE."),
    ]


def test_builds_faq_payload_preserving_section_settings_and_dropdown_ids():
    entries = extract_common_questions("pdp.docx", _faq_doc())
    existing = {
        "id": 9,
        "dropdownInitialOpen": False,
        "idForScrolling": "commonQuestions",
        "hideSection": False,
        "title": {"desktop": "Preguntas frecuentes"},
        "dropdowns": [
            {"id": 10, "dropdownTitle": "Texto anterior", "richText": "Respuesta anterior", "iconPosition": "Right"},
            {"id": 11, "dropdownTitle": "Otra pregunta", "richText": "Otra respuesta", "iconPosition": "Left"},
        ],
    }
    assert not common_questions_match(entries, existing)
    payload = build_common_questions_payload(entries, existing)
    assert payload["id"] == 9
    assert payload["dropdownInitialOpen"] is False
    assert payload["dropdowns"] == [
        {"id": 10, "iconPosition": "Right", "dropdownTitle": "¿Cuánto dura el programa?", "richText": "Dura 48 meses."},
        {"id": 11, "iconPosition": "Left", "dropdownTitle": "¿Tiene validez oficial?", "richText": "Sí, cuenta con RVOE."},
    ]
    assert common_questions_match(entries, payload)


def test_description_runner_dry_run_reports_faq_mismatch_and_live_run_syncs(tmp_path):
    doc = Document()
    doc.add_paragraph("Programa A", style="Title")
    doc.add_paragraph("Descripción del programa.")
    doc.add_paragraph("Asignaturas")
    doc.add_paragraph("Contenido del plan.")
    doc.add_paragraph("1° cuatrimestre")
    doc.add_heading("12. Preguntas frecuentes (FAQ)", level=2)
    doc.add_paragraph("1. ¿Cuánto dura el programa?")
    doc.add_paragraph("Dura 48 meses.")
    path = tmp_path / "programa.docx"
    doc.save(path)

    class FakeClient:
        def __init__(self):
            self.updates = []

        async def find_product(self, *args, **kwargs): return {"id": 12}
        async def get_product_common_questions(self, *args):
            return {"id": 9, "dropdownInitialOpen": True, "dropdowns": [
                {"id": 10, "dropdownTitle": "¿Cuánto dura el programa?", "richText": "Dura 44 meses.", "iconPosition": "Left"},
            ]}
        async def update_product(self, identifier, attributes): self.updates.append((identifier, attributes))

    client = FakeClient()
    row = ProductRow("Sheet", 2, "Programa A", str(path))
    dry_results, _ = asyncio.run(StrapiDescriptionRunner(client, "Argentina", "es-AR", dry_run=True).run([row]))
    assert dry_results[0].status == "DRY_RUN"
    assert "Preguntas frecuentes se actualizarían" in dry_results[0].message
    assert client.updates == []

    results, _ = asyncio.run(StrapiDescriptionRunner(client, "Argentina", "es-AR", dry_run=False).run([row]))
    assert results[0].status == "UPDATED"
    faq_payload = client.updates[-1][1]["commonQuestions"]
    assert faq_payload["dropdowns"][0]["dropdownTitle"] == "¿Cuánto dura el programa?"
    assert faq_payload["dropdowns"][0]["richText"] == "Dura 48 meses."


def test_description_runner_does_not_rewrite_common_questions_when_already_equal(tmp_path):
    doc = Document()
    doc.add_paragraph("Programa A", style="Title")
    doc.add_paragraph("Descripción del programa.")
    doc.add_paragraph("Asignaturas")
    doc.add_paragraph("Contenido del plan.")
    doc.add_paragraph("1° cuatrimestre")
    doc.add_heading("12. Preguntas frecuentes (FAQ)", level=2)
    doc.add_paragraph("1. ¿Cuánto dura el programa?")
    doc.add_paragraph("Dura 48 meses.")
    path = tmp_path / "programa.docx"
    doc.save(path)
    expected = extract_common_questions("programa.docx", path.read_bytes())
    existing = build_common_questions_payload(expected, {"id": 9, "dropdowns": [
        {"id": 10, "dropdownTitle": "Anterior", "richText": "Anterior", "iconPosition": "Left"},
    ]})

    class FakeClient:
        def __init__(self): self.updates = []
        async def find_product(self, *args, **kwargs): return {"id": 12}
        async def get_product_common_questions(self, *args): return existing
        async def update_product(self, identifier, attributes): self.updates.append(attributes)

    client = FakeClient()
    result, _ = asyncio.run(StrapiDescriptionRunner(client, "Argentina", "es-AR", dry_run=False).run([
        ProductRow("Sheet", 2, "Programa A", str(path)),
    ]))
    assert result[0].status == "UPDATED"
    assert "commonQuestions" not in client.updates[0]
