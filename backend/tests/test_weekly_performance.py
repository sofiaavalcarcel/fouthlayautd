"""Pruebas del procesamiento Excel + PageSpeed de Weekly Performance."""

import asyncio
import io
import time
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app.config.settings import Settings
from backend.app.main import create_app
from backend.app.modules.weekly_auto.weekly_performance import (
    WeeklyPerformanceConfig,
    WeeklyPerformanceError,
    WeeklyPerformanceRunner,
)


def _workbook_bytes(*urls: str, sheet_name: str = "Hoja 1") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.cell(1, 4, "URL")
    sheet.cell(1, 12, "Desktop")
    sheet.cell(1, 13, "Mobile")
    for row, url in enumerate(urls, start=2):
        sheet.cell(row, 4, url)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _unsized_workbook_bytes(*urls: str) -> bytes:
    """Quita la dimensión XML para reproducir exportaciones sin max_row."""

    source = _workbook_bytes(*urls)
    output = io.BytesIO()
    with ZipFile(io.BytesIO(source), "r") as workbook_zip, ZipFile(
        output, "w", compression=ZIP_DEFLATED
    ) as patched_zip:
        for entry in workbook_zip.infolist():
            content = workbook_zip.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(
                    b'<dimension ref="D1:M2"/>', b""
                )
            patched_zip.writestr(entry, content)
    return output.getvalue()


def test_runner_writes_desktop_and_mobile_scores_and_preserves_workbook(tmp_path, monkeypatch):
    content = _workbook_bytes("https://example.com/a", "not-a-url", "https://example.com/b")
    runner = WeeklyPerformanceRunner(Settings(storage_dir=tmp_path / "storage"))
    calls: list[tuple[str, str]] = []

    async def fake_score(client, url, strategy, api_key, should_stop):
        calls.append((url, strategy))
        return 91 if strategy == "desktop" else 73

    monkeypatch.setattr(runner, "_score", fake_score)
    output_path = tmp_path / "resultados.xlsx"
    progress = []
    result = asyncio.run(
        runner.run(
            content,
            output_path,
            WeeklyPerformanceConfig(),
            progress_callback=progress.append,
        )
    )

    workbook = load_workbook(output_path, data_only=True)
    sheet = workbook["Hoja 1"]
    assert (sheet.cell(2, 12).value, sheet.cell(2, 13).value) == (91, 73)
    assert (sheet.cell(4, 12).value, sheet.cell(4, 13).value) == (91, 73)
    assert sheet.cell(3, 12).value is None
    workbook.close()
    assert result["status"] == "PASS"
    assert result["total_urls"] == result["completed"] == 2
    assert len(progress) == 2
    assert len(calls) == 4


def test_runner_records_pagespeed_failures_in_the_output(tmp_path, monkeypatch):
    runner = WeeklyPerformanceRunner(Settings(storage_dir=tmp_path / "storage"))

    async def fake_score(client, url, strategy, api_key, should_stop):
        return "Error 429" if strategy == "mobile" else 88

    monkeypatch.setattr(runner, "_score", fake_score)
    output_path = tmp_path / "with-error.xlsx"
    result = asyncio.run(
        runner.run(_workbook_bytes("https://example.com"), output_path, WeeklyPerformanceConfig())
    )
    workbook = load_workbook(output_path, data_only=True)
    sheet = workbook["Hoja 1"]
    assert sheet.cell(2, 12).value == 88
    assert sheet.cell(2, 13).value == "Error 429"
    workbook.close()
    assert result["status"] == "FAIL"
    assert result["failed"] == 1


def test_inspect_rejects_missing_sheet_and_workbook_without_urls(tmp_path):
    runner = WeeklyPerformanceRunner(Settings(storage_dir=tmp_path / "storage"))
    with pytest.raises(WeeklyPerformanceError, match="No existe la pestaña"):
        runner.inspect(_workbook_bytes("https://example.com", sheet_name="Datos"), "Hoja 1")
    with pytest.raises(WeeklyPerformanceError, match="columna D"):
        runner.inspect(_workbook_bytes("texto"), "Hoja 1")


def test_inspect_accepts_excel_without_dimension_metadata(tmp_path):
    """La columna D se puede leer aunque la hoja no declare su dimensión."""

    runner = WeeklyPerformanceRunner(Settings(storage_dir=tmp_path / "storage"))
    inspected = runner.inspect(_unsized_workbook_bytes("https://example.com"), "Hoja 1")
    assert inspected["rows"] == [2]


def test_pagespeed_key_is_only_loaded_from_settings(tmp_path):
    settings = Settings(storage_dir=tmp_path / "storage", pagespeed_api_key="private-test-key")
    runner = WeeklyPerformanceRunner(settings)
    assert runner.settings.pagespeed_api_key.get_secret_value() == "private-test-key"


def test_api_runs_reports_status_and_downloads_result(tmp_path, monkeypatch):
    async def fake_run(self, content, output_path, config, *, should_stop=None, progress_callback=None):
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(content)
        progress_callback(
            {
                "row": 2,
                "url": "https://example.com",
                "desktop": 90,
                "mobile": 75,
                "status": "PASS",
                "completed": 1,
                "total": 1,
            }
        )
        return {
            "status": "PASS",
            "summary": "1 URLs procesadas; 0 presentaron error de PageSpeed.",
            "completed": 1,
            "successful": 1,
            "failed": 0,
            "results": [],
        }

    monkeypatch.setattr(WeeklyPerformanceRunner, "run", fake_run)
    app = create_app(
        Settings(database_path=tmp_path / "qa.db", storage_dir=tmp_path / "storage")
    )
    with TestClient(app) as client:
        response = client.post(
            "/api/weekly-auto/performance/run",
            files={
                "file": (
                    "performance.xlsx",
                    _workbook_bytes("https://example.com"),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"sheet_name": "Hoja 1", "max_workers": "4"},
        )
        assert response.status_code == 202
        job_id = response.json()["job_id"]
        job = response.json()
        for _ in range(50):
            job = client.get(f"/api/weekly-auto/performance/runs/{job_id}").json()
            if job["status"] not in {"QUEUED", "RUNNING"}:
                break
            time.sleep(0.01)
        assert job["status"] == "PASS"
        assert job["successful"] == 1
        assert "_output_path" not in job
        download = client.get(f"/api/weekly-auto/performance/runs/{job_id}/download")
        assert download.status_code == 200
        assert download.content.startswith(b"PK")
