from __future__ import annotations

import re
import unicodedata
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


TAB_PREFIXES = ("Perfil ingreso", "Perfil egreso", "Empleabilidad")
_ALIASES = {
    "Perfil ingreso": ("perfil ingreso", "perfil de ingreso"),
    "Perfil egreso": ("perfil egreso", "perfil de egreso"),
    "Empleabilidad": ("empleabilidad", "donde podras trabajar", "campo laboral", "salida laboral", "oportunidades profesionales"),
}
_SECTION_BOUNDARIES = {
    "validez academica",
    "asignaturas",
    "areas de concentracion",
    "elige en que modalidad estudiar",
    "elige una licenciatura",
}


@dataclass(frozen=True)
class BulletTabContent:
    prefix: str
    description: str
    bullets: tuple[str, ...]


def extract_study_modalities(filename: str, content: bytes) -> tuple[str, ...]:
    """Read the modality labels from the FT's study-modality table."""
    extension = Path(urlparse(filename).path or filename).suffix.casefold()
    if extension != ".docx":
        return ()
    try:
        document = Document(BytesIO(content))
    except Exception as error:
        raise ValueError(f"No se pudo leer las modalidades de estudio en {filename}.") from error
    found: list[str] = []
    for table in document.tables:
        for row in table.rows:
            labels = [_clean(cell.text.splitlines()[0]) if cell.text.splitlines() else "" for cell in row.cells]
            normalized = [_normalize(label) for label in labels]
            if any(value in {"en linea", "ejecutiva", "hibrida"} for value in normalized):
                for label, value in zip(labels, normalized):
                    if value in {"en linea", "ejecutiva", "hibrida"} and label not in found:
                        found.append(label)
                return tuple(found)
    return ()


def _normalize(text: str) -> str:
    plain = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z0-9 ]+", " ", plain.casefold()).split())


def _clean(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").split()).strip()


def _is_bullet(paragraph) -> bool:
    style = (paragraph.style.name if paragraph.style else "").casefold()
    if "bullet" in style or "list paragraph" in style:
        return True
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.find(qn("w:numPr")) is not None


def _section_prefix(text: str) -> str | None:
    value = _normalize(text)
    for prefix, aliases in _ALIASES.items():
        if value in {_normalize(alias) for alias in aliases}:
            return prefix
    return None


def _read_docx_paragraphs(content: bytes) -> list[tuple[str, bool]]:
    document = Document(BytesIO(content))
    return [(_clean(paragraph.text), _is_bullet(paragraph)) for paragraph in document.paragraphs]


def _extract_profile_table_sections(document: Document) -> dict[str, BulletTabContent]:
    """Extract ingreso/egreso from the PDP's Perfil del estudiante table."""
    sections: dict[str, BulletTabContent] = {}
    for table in document.tables:
        if not table.rows or len(table.rows[0].cells) < 2:
            continue
        for cell in table.rows[0].cells:
            paragraphs = [_clean(paragraph.text) for paragraph in cell.paragraphs if _clean(paragraph.text)]
            if not paragraphs:
                continue
            prefix = _section_prefix(paragraphs[0])
            if prefix not in {"Perfil ingreso", "Perfil egreso"}:
                continue
            bullets = tuple(paragraphs[1:])
            if bullets:
                # The table heading identifies the profile; it is not an
                # introductory description for bulletsDescription.
                sections[prefix] = BulletTabContent(prefix, "", bullets)
    return sections


def _read_pdf_paragraphs(content: bytes) -> list[tuple[str, bool]]:
    lines = [line.strip() for page in PdfReader(BytesIO(content)).pages for line in (page.extract_text() or "").splitlines()]
    return [(_clean(line), bool(re.match(r"^[•·▪◦\-*]\s*", line))) for line in lines if _clean(line)]


def extract_bullet_tabs(filename: str, content: bytes) -> dict[str, BulletTabContent]:
    """Extract introductory copy and actual list items for each PDP bullet tab."""
    extension = Path(urlparse(filename).path or filename).suffix.casefold()
    try:
        if extension == ".docx":
            paragraphs = _read_docx_paragraphs(content)
        elif extension == ".pdf":
            paragraphs = _read_pdf_paragraphs(content)
        else:
            raise ValueError(f"Formato de Documento PDP no soportado para Bullet Tabs: {filename}.")
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"No se pudo leer el contenido de Bullet Tabs en {filename}.") from error

    sections: dict[str, BulletTabContent] = {}
    index = 0
    while index < len(paragraphs):
        prefix = _section_prefix(paragraphs[index][0])
        if prefix is None:
            index += 1
            continue
        index += 1
        intro: list[str] = []
        bullets: list[str] = []
        found_bullet = False
        while index < len(paragraphs):
            text, is_bullet = paragraphs[index]
            if _section_prefix(text):
                break
            if _normalize(text).rstrip("*") in _SECTION_BOUNDARIES:
                break
            index += 1
            if not text:
                continue
            if is_bullet:
                found_bullet = True
                cleaned = re.sub(r"^[•·▪◦\-*]\s*", "", text).strip()
                if cleaned:
                    bullets.append(cleaned)
            elif found_bullet:
                # Other headings and following prose belong outside this tab's bullets.
                continue
            else:
                intro.append(text)
        # Some FT versions use plain paragraphs instead of Word bullet styles.
        # In those documents the section shape is stable: ingreso has an intro
        # followed by items, egreso is only a list, and employability is made
        # of area/description pairs.
        if not found_bullet and intro:
            if prefix == "Perfil ingreso":
                intro, bullets = intro[:1], intro[1:]
            elif prefix == "Perfil egreso":
                bullets = intro
                intro = ["Perfil de egreso"]
            elif prefix == "Empleabilidad":
                pairs = []
                for pair_index in range(0, len(intro), 2):
                    title = intro[pair_index]
                    description = intro[pair_index + 1] if pair_index + 1 < len(intro) else ""
                    pairs.append(f"{title}: {description}".strip(": "))
                bullets = pairs
                intro = ["Dónde podrás trabajar"]
        if not intro and bullets and prefix == "Perfil egreso":
            intro = ["Perfil de egreso"]
        if not intro or not bullets:
            raise ValueError(
                f"La sección {prefix!r} de {filename} debe tener un párrafo introductorio y al menos una viñeta."
            )
        if prefix in sections:
            raise ValueError(f"La sección {prefix!r} aparece más de una vez en {filename}.")
        sections[prefix] = BulletTabContent(prefix, "\n\n".join(intro), tuple(bullets))

    if extension == ".docx":
        document = Document(BytesIO(content))
        # The PDP table is authoritative for these two profiles. Employment
        # remains extracted from its existing section above.
        sections.update(_extract_profile_table_sections(document))

    # Missing sections are intentionally returned empty. The caller preserves
    # the existing Strapi tab and does not create or overwrite blank content.
    return sections


def _tab_title(prefix: str, program: str) -> str:
    return f"{prefix} {program}"


def linked_tab_prefix(strapi_name: str, program: str) -> str | None:
    """Identify a product's tab from its relation name, accepting the optional 'de'."""
    value = _normalize(strapi_name)
    program_name = _normalize(program)
    for prefix, aliases in _ALIASES.items():
        for alias in aliases:
            if value == f"{_normalize(alias)} {program_name}":
                return prefix
    return None


def _bullets_component(entry: dict, template: dict | None, section: BulletTabContent, *, creating: bool) -> dict:
    content_items = (entry.get("attributes", entry).get("content") or []) if entry else []
    current = next((item for item in content_items if item.get("__component") == "section.bullets"), None)
    source = current or next(
        (item for item in ((template or {}).get("attributes", template or {}).get("content") or []) if item.get("__component") == "section.bullets"),
        {},
    )
    component = deepcopy(source)
    if creating:
        _remove_component_ids(component)
    template_component = next(
        (item for item in ((template or {}).get("attributes", template or {}).get("content") or []) if item.get("__component") == "section.bullets"),
        {},
    )
    _copy_cover_image(component, template_component, creating=creating)
    component["__component"] = "section.bullets"
    if section.description:
        component["bulletsDescription"] = _desktop_value(component.get("bulletsDescription"), section.description)
    else:
        # Explicit null clears an existing Strapi component instead of
        # allowing an omitted field to survive the update.
        component["bulletsDescription"] = None
    component["bullets"] = [
        {
            "title": None,
            "description": text,
            "chakraConfig": None,
            "icon": {"name": "UilCheck", "chakraConfig": None},
            "iconMobile": None,
        }
        for text in section.bullets
    ]
    return component


def has_desktop_cover_image(entry: dict | None) -> bool:
    """Return whether a bullet tab has an image assigned for desktop."""
    attributes = (entry or {}).get("attributes", entry or {})
    component = next((item for item in attributes.get("content") or [] if item.get("__component") == "section.bullets"), {})
    cover = component.get("coverImage") or {}
    desktop = cover.get("desktop") or {}
    image = desktop.get("image") or {}
    if isinstance(image, dict) and isinstance(image.get("data"), dict):
        image = image["data"]
    return isinstance(image, dict) and image.get("id") is not None


def _copy_cover_image(component: dict, template_component: dict, *, creating: bool) -> None:
    """Use the same-locale tab's desktop image when this tab has none."""
    current_cover = deepcopy(component.get("coverImage") or {})
    template_cover = deepcopy(template_component.get("coverImage") or {})
    current_desktop = current_cover.get("desktop") or {}
    current_image = current_desktop.get("image") or {}
    if isinstance(current_image, dict) and isinstance(current_image.get("data"), dict):
        current_image = current_image["data"]
    has_current_image = isinstance(current_image, dict) and current_image.get("id") is not None
    if has_current_image:
        # Populate responses wrap media in `data`; Strapi updates expect a
        # relation reference with only the media ID.
        current_desktop["image"] = {"id": current_image["id"]}
        current_cover["desktop"] = current_desktop
        component["coverImage"] = current_cover
        return

    template_desktop = deepcopy(template_cover.get("desktop") or {})
    image = template_desktop.get("image") or {}
    if isinstance(image, dict) and isinstance(image.get("data"), dict):
        image = image["data"]
    if not isinstance(image, dict) or image.get("id") is None:
        return

    # The desktop/media components belong to the reference tab. Reuse only its
    # media asset and settings, never its component IDs.
    template_desktop.pop("id", None)
    template_desktop["image"] = {"id": image["id"]}
    if creating:
        for key in ("mobile", "tablet"):
            device = template_cover.get(key)
            if isinstance(device, dict):
                device = deepcopy(device)
                device.pop("id", None)
                image_value = device.get("image") or {}
                if isinstance(image_value, dict) and isinstance(image_value.get("data"), dict):
                    image_value = image_value["data"]
                if isinstance(image_value, dict) and image_value.get("id") is not None:
                    device["image"] = {"id": image_value["id"]}
                template_cover[key] = device
        template_cover.pop("id", None)
    elif current_cover.get("id") is not None:
        template_cover["id"] = current_cover["id"]
    template_cover["desktop"] = template_desktop
    component["coverImage"] = template_cover


def _remove_component_ids(value: dict | list) -> None:
    """Remove nested Strapi component IDs while retaining media relation IDs."""
    if isinstance(value, list):
        for item in value:
            if isinstance(item, (dict, list)):
                _remove_component_ids(item)
        return
    if value.get("__component") or "chakraConfig" in value:
        value.pop("id", None)
    for key, item in value.items():
        if key == "coverImage":
            continue
        if isinstance(item, (dict, list)):
            _remove_component_ids(item)


def _desktop_value(current: dict | str | None, value: str) -> dict:
    result = deepcopy(current) if isinstance(current, dict) else {}
    result.pop("id", None)
    result["desktop"] = value
    result.setdefault("mobile", None)
    result.setdefault("chakraConfig", None)
    return result


def build_bullet_tab_payload(
    program: str,
    section: BulletTabContent,
    *,
    existing: dict | None = None,
    template: dict | None = None,
    locale: str,
) -> dict:
    """Build a tab payload, preserving its own visual settings and never the template's IDs."""
    entry = existing or {}
    attributes = entry.get("attributes", entry)
    template_attributes = (template or {}).get("attributes", template or {})
    content_template = template_attributes.get("content") or []
    content_items = deepcopy(attributes.get("content") or content_template)
    if not existing:
        _remove_component_ids(content_items)
    old_bullets_component = next(
        (item for item in content_items if item.get("__component") == "section.bullets"),
        None,
    )
    new_component = _bullets_component(entry, template, section, creating=not bool(existing))
    if old_bullets_component is not None:
        target_index = content_items.index(old_bullets_component)
        content_items[target_index] = new_component
    elif content_items:
        content_items.append(new_component)
    else:
        content_items = [new_component]

    # Keep only the schema-backed fields needed to create a new record. Existing
    # entries retain their own title/configuration while only their tab content changes.
    if existing:
        return {"content": content_items}

    title_source = deepcopy(template_attributes.get("title") or {})
    if isinstance(title_source, dict):
        title_source.pop("id", None)
        title_source["desktop"] = section.prefix
        title_source.setdefault("mobile", None)
        title_source.setdefault("chakraConfig", None)
    return {
        "strapiName": _tab_title(section.prefix, program),
        "locale": locale,
        "title": title_source or {"desktop": section.prefix, "mobile": None, "chakraConfig": None},
        "iconTitle": deepcopy(template_attributes.get("iconTitle")),
        "content": content_items,
    }
