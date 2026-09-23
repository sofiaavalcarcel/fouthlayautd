from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from .models import ProductRow


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    # Algunos Excel exportados desde Google Sheets conservan el apostrofo
    # usado para forzar texto en la celda.
    if text.startswith("'"):
        text = text[1:].strip()
    return text or None


def _hyperlink(cell: Any) -> str | None:
    """Return the real URL stored in an Excel/Google Sheets hyperlink."""
    link = getattr(cell, "hyperlink", None)
    target = getattr(link, "target", None) if link else None
    return _text(target)


def read_product_rows(content: bytes) -> list[ProductRow]:
    # read_only mode omits hyperlink metadata in some XLSX files exported by
    # Google Sheets, so load the workbook normally to preserve those URLs.
    workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    rows: list[ProductRow] = []
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows()
            header_row_number = None
            header_row = None
            for candidate_number, candidate in enumerate(iterator, start=1):
                candidate_headers = {
                    _text(cell.value).casefold(): index
                    for index, cell in enumerate(candidate)
                    if _text(cell.value)
                }
                if any(name in candidate_headers for name in ("programa", "plan de estudios")):
                    header_row_number = candidate_number
                    header_row = candidate
                    break

            # The new source workbook contains auxiliary QA and reference
            # sheets. Only sheets with a product-name column are input sheets.
            if header_row is None or header_row_number is None:
                continue

            headers = {_text(cell.value).casefold(): index for index, cell in enumerate(header_row) if _text(cell.value)}

            program_index = next(headers[name] for name in ("programa", "plan de estudios") if name in headers)
            document_index = next(
                (headers[name] for name in ("documento pdp", "documento", "pdp", "referencia") if name in headers),
                None,
            )
            ft_index = headers.get("ft")
            # Validation matrices can contain a later row labelled
            # "Plan de estudios" as a checklist item, but they have no PDP
            # source column and are not product input sheets.
            if "plan de estudios" in headers and document_index is None:
                continue
            for row_number, row in enumerate(iterator, start=header_row_number + 1):
                program = _text(row[program_index].value if program_index < len(row) else None)
                if program:
                    document = None
                    if document_index is not None and document_index < len(row):
                        # The visible cell value is normally only the PDP name.
                        # Prefer its hyperlink, which points to the actual Google Doc.
                        document = _hyperlink(row[document_index]) or _text(row[document_index].value)
                    ft_document = None
                    if ft_index is not None and ft_index < len(row):
                        ft_document = _hyperlink(row[ft_index]) or _text(row[ft_index].value)
                    rows.append(ProductRow(sheet.title, row_number, program, document, ft_document))
        if not rows:
            raise ValueError("El Excel no contiene programas para procesar ni una columna Programa/Plan de estudios válida.")
        return rows
    finally:
        workbook.close()


def select_product_rows(rows: list[ProductRow], product_scope: str) -> list[ProductRow]:
    """Apply an all or first-N scope while preserving spreadsheet order."""
    if product_scope == "all":
        return rows
    try:
        count = int(product_scope)
    except (TypeError, ValueError) as error:
        raise ValueError("product_scope debe ser un entero positivo o all.") from error
    if count < 1:
        raise ValueError("product_scope debe ser un entero positivo o all.")
    return rows[:count]
