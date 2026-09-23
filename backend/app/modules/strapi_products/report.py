from __future__ import annotations

import io
import json

from openpyxl import Workbook

from .models import ProductResult, ProductSummary


def build_report(results: list[ProductResult], summary: ProductSummary) -> bytes:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Resultados"
    detail.append(["Hoja", "Fila", "Programa", "Pais", "Estado", "Canonical anterior", "Canonical nuevo", "Mensaje"])
    for result in results:
        detail.append([result.sheet, result.row_number, result.program, result.country, result.status, result.old_canonical, result.new_canonical, result.message])

    overview = workbook.create_sheet("Resumen")
    overview.append(["Metrica", "Valor"])
    for key, value in summary.as_dict().items():
        overview.append([key, value])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_description_report(results: list[ProductResult], summary: ProductSummary) -> bytes:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Resultados"
    detail.append(["Hoja", "Fila", "Programa", "Pais", "Estado", "siuKey", "bannerKey", "Descripcion extraida", "Cambios", "Verificacion", "Mensaje"])
    for result in results:
        detail.append([
            result.sheet, result.row_number, result.program, result.country, result.status,
            result.siu_key, result.banner_key, result.description, len(result.changes),
            result.verification.get("ok"), result.message,
        ])
    overview = workbook.create_sheet("Resumen")
    overview.append(["Metrica", "Valor"])
    for key, value in summary.as_dict().items():
        overview.append([key, value])
    overview.append(["cambios_propuestos_o_aplicados", sum(len(result.changes) for result in results)])
    overview.append(["productos_verificados", sum(result.verification.get("ok") is True for result in results)])

    plan = workbook.create_sheet("Plan por producto")
    plan.append(["Hoja", "Fila", "Programa", "Estado", "Campo", "Accion", "Antes", "Despues", "Verificado", "ID creado"])
    for result in results:
        for change in result.changes:
            before = json.dumps(change.get("before"), ensure_ascii=False, default=str) if not isinstance(change.get("before"), str) else change.get("before")
            after = json.dumps(change.get("after"), ensure_ascii=False, default=str) if not isinstance(change.get("after"), str) else change.get("after")
            plan.append([
                result.sheet, result.row_number, result.program, result.status,
                change.get("field"), change.get("action"), before, after,
                change.get("verified"), change.get("created_id"),
            ])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_combined_report(canonical_job: dict, pdp_job: dict) -> bytes:
    """Build one workbook with both stages of a PDP product synchronization."""
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Resultados"
    detail.append([
        "Proceso", "Hoja", "Fila", "Programa", "Pais", "Estado", "Canonical anterior",
        "Canonical nuevo", "siuKey", "bannerKey", "Descripcion extraida", "Verificacion", "Mensaje",
    ])
    all_results = [("Canonical", item) for item in canonical_job.get("results", [])]
    all_results += [("PDP", item) for item in pdp_job.get("results", [])]
    for process, result in all_results:
        detail.append([
            process, result.get("sheet"), result.get("row_number"), result.get("program"),
            result.get("country"), result.get("status"), result.get("old_canonical"),
            result.get("new_canonical"), result.get("siu_key"), result.get("banner_key"),
            result.get("description"), result.get("verification", {}).get("ok"), result.get("message"),
        ])

    overview = workbook.create_sheet("Resumen")
    overview.append(["Proceso", "Estado", "Metrica", "Valor"])
    for process, job in (("Canonical", canonical_job), ("PDP", pdp_job)):
        summary = job.get("summary", {})
        overview.append([process, job.get("status"), "total", summary.get("total", 0)])
        for key, value in summary.items():
            if key != "total":
                overview.append([process, job.get("status"), key, value])

    plan = workbook.create_sheet("Plan por producto")
    plan.append(["Proceso", "Hoja", "Fila", "Programa", "Estado", "Campo", "Accion", "Antes", "Despues", "Verificado", "ID creado"])
    for process, result in all_results:
        for change in result.get("changes", []):
            before = json.dumps(change.get("before"), ensure_ascii=False, default=str)
            after = json.dumps(change.get("after"), ensure_ascii=False, default=str)
            plan.append([
                process, result.get("sheet"), result.get("row_number"), result.get("program"),
                result.get("status"), change.get("field"), change.get("action"), before, after,
                change.get("verified"), change.get("created_id"),
            ])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
