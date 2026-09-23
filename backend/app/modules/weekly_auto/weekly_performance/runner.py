"""PageSpeed Insights concurrente para el módulo Weekly Performance."""

from __future__ import annotations

import asyncio
import io
from datetime import datetime
from pathlib import Path
from time import perf_counter
from typing import Any, Callable

import httpx
from openpyxl import load_workbook

from ....config.settings import Settings
from ....services.logging_service import get_logger
from .schemas import WeeklyPerformanceConfig


class WeeklyPerformanceError(RuntimeError):
    """Error legible del procesamiento de rendimiento."""


class WeeklyPerformanceCancelled(RuntimeError):
    """La ejecución fue detenida por el usuario."""


class WeeklyPerformanceRunner:
    """Lee D, consulta PageSpeed y escribe desktop/móvil en L/M."""

    ENDPOINT = "https://www.googleapis.com/pagespeedonline/v5/runPagespeed"
    URL_COLUMN = 4
    DESKTOP_COLUMN = 12
    MOBILE_COLUMN = 13
    FIRST_ROW = 2
    MAX_RETRIES = 3
    BACKOFF_BASE_SECONDS = 4
    CHECKPOINT_EVERY = 10

    def __init__(self, settings: Settings):
        self.settings = settings
        self.logger = get_logger()

    @staticmethod
    def _valid_url(value: Any) -> str:
        url = str(value or "").strip()
        return url if url.casefold().startswith(("http://", "https://")) else ""

    def inspect(self, content: bytes, sheet_name: str) -> dict[str, Any]:
        """Valida el libro y cuenta las URLs antes de crear el trabajo."""

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except Exception as error:  # noqa: BLE001 - openpyxl expone varios tipos
            raise WeeklyPerformanceError(f"No se pudo abrir el Excel: {error}") from error
        try:
            if sheet_name not in workbook.sheetnames:
                available = ", ".join(workbook.sheetnames)
                raise WeeklyPerformanceError(
                    f"No existe la pestaña '{sheet_name}'. Disponibles: {available}."
                )
            worksheet = workbook[sheet_name]
            # Algunos Excel exportados por Google Sheets u otras herramientas
            # omiten <dimension> en el XML. En modo read_only, openpyxl deja
            # ``max_row`` en None para esas hojas; iterar la columna evita el
            # TypeError y además conserva el bajo consumo de memoria.
            rows = [
                row_number
                for row_number, (value,) in enumerate(
                    worksheet.iter_rows(
                        min_row=self.FIRST_ROW,
                        min_col=self.URL_COLUMN,
                        max_col=self.URL_COLUMN,
                        values_only=True,
                    ),
                    start=self.FIRST_ROW,
                )
                if self._valid_url(value)
            ]
            if not rows:
                raise WeeklyPerformanceError(
                    "No se encontraron URLs http/https en la columna D."
                )
            return {"sheet_name": sheet_name, "total_urls": len(rows), "rows": rows}
        finally:
            workbook.close()

    async def _interruptible_sleep(
        self,
        seconds: float,
        should_stop: Callable[[], bool] | None,
    ) -> None:
        deadline = perf_counter() + seconds
        while perf_counter() < deadline:
            if should_stop and should_stop():
                raise WeeklyPerformanceCancelled("Ejecución detenida por el usuario.")
            await asyncio.sleep(min(0.25, max(0.0, deadline - perf_counter())))

    async def _score(
        self,
        client: httpx.AsyncClient,
        url: str,
        strategy: str,
        api_key: str,
        should_stop: Callable[[], bool] | None,
    ) -> int | str:
        params = {"url": url, "strategy": strategy, "category": "performance"}
        if api_key:
            params["key"] = api_key

        for attempt in range(self.MAX_RETRIES + 1):
            if should_stop and should_stop():
                raise WeeklyPerformanceCancelled("Ejecución detenida por el usuario.")
            try:
                request_task = asyncio.create_task(client.get(self.ENDPOINT, params=params))
                while not request_task.done():
                    if should_stop and should_stop():
                        request_task.cancel()
                        await asyncio.gather(request_task, return_exceptions=True)
                        raise WeeklyPerformanceCancelled("Ejecución detenida por el usuario.")
                    await asyncio.wait({request_task}, timeout=0.25)
                response = await request_task
                if response.status_code == 429:
                    if attempt < self.MAX_RETRIES:
                        retry_after = response.headers.get("Retry-After", "")
                        wait = (
                            int(retry_after)
                            if retry_after.isdigit()
                            else self.BACKOFF_BASE_SECONDS * (2**attempt)
                        )
                        await self._interruptible_sleep(wait, should_stop)
                        continue
                    return "Error 429"
                response.raise_for_status()
                payload = response.json()
                score = payload["lighthouseResult"]["categories"]["performance"]["score"]
                return int(score * 100) if score is not None else "N/A"
            except (httpx.RequestError, httpx.HTTPStatusError, ValueError, KeyError, TypeError) as error:
                if attempt < self.MAX_RETRIES:
                    await self._interruptible_sleep(
                        self.BACKOFF_BASE_SECONDS * (2**attempt), should_stop
                    )
                    continue
                if isinstance(error, httpx.HTTPStatusError):
                    return f"Error HTTP {error.response.status_code}"
                return f"Error: {str(error)[:180]}"
        return "Error desconocido"

    async def _process_row(
        self,
        semaphore: asyncio.Semaphore,
        client: httpx.AsyncClient,
        row: int,
        url: str,
        api_key: str,
        should_stop: Callable[[], bool] | None,
    ) -> dict[str, Any]:
        async with semaphore:
            started = perf_counter()
            mobile = await self._score(client, url, "mobile", api_key, should_stop)
            desktop = await self._score(client, url, "desktop", api_key, should_stop)
            status = "PASS" if isinstance(mobile, int) and isinstance(desktop, int) else "FAIL"
            return {
                "row": row,
                "url": url,
                "desktop": desktop,
                "mobile": mobile,
                "status": status,
                "elapsed_seconds": round(perf_counter() - started, 2),
            }

    async def run(
        self,
        content: bytes,
        output_path: Path,
        config: WeeklyPerformanceConfig,
        *,
        should_stop: Callable[[], bool] | None = None,
        progress_callback: Callable[[dict[str, Any]], None] | None = None,
    ) -> dict[str, Any]:
        """Ejecuta las consultas concurrentes y conserva checkpoints parciales."""

        inspected = self.inspect(content, config.sheet_name)
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=False)
        worksheet = workbook[config.sheet_name]
        output_path.parent.mkdir(parents=True, exist_ok=True)
        tasks: list[asyncio.Task] = []
        results: list[dict[str, Any]] = []
        started_at = datetime.now().isoformat(timespec="seconds")
        started = perf_counter()
        api_key = self.settings.pagespeed_api_key.get_secret_value().strip()
        semaphore = asyncio.Semaphore(config.max_workers)

        try:
            timeout = httpx.Timeout(120.0, connect=30.0)
            async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as client:
                for row in inspected["rows"]:
                    url = self._valid_url(worksheet.cell(row=row, column=self.URL_COLUMN).value)
                    tasks.append(
                        asyncio.create_task(
                            self._process_row(
                                semaphore, client, row, url, api_key, should_stop
                            )
                        )
                    )

                for future in asyncio.as_completed(tasks):
                    if should_stop and should_stop():
                        raise WeeklyPerformanceCancelled("Ejecución detenida por el usuario.")
                    result = await future
                    results.append(result)
                    worksheet.cell(
                        row=result["row"], column=self.DESKTOP_COLUMN, value=result["desktop"]
                    )
                    worksheet.cell(
                        row=result["row"], column=self.MOBILE_COLUMN, value=result["mobile"]
                    )
                    if progress_callback:
                        progress_callback(
                            {
                                **result,
                                "completed": len(results),
                                "total": inspected["total_urls"],
                            }
                        )
                    if len(results) % self.CHECKPOINT_EVERY == 0:
                        await asyncio.to_thread(workbook.save, output_path)
        finally:
            for task in tasks:
                if not task.done():
                    task.cancel()
            if tasks:
                await asyncio.gather(*tasks, return_exceptions=True)
            # Incluso una cancelación deja un Excel parcial recuperable.
            await asyncio.shield(asyncio.to_thread(workbook.save, output_path))
            workbook.close()

        failed = sum(result["status"] == "FAIL" for result in results)
        finished_at = datetime.now().isoformat(timespec="seconds")
        return {
            "status": "PASS" if failed == 0 else "FAIL",
            "summary": f"{len(results)} URLs procesadas; {failed} presentaron error de PageSpeed.",
            "started_at": started_at,
            "finished_at": finished_at,
            "duration_seconds": round(perf_counter() - started, 2),
            "total_urls": inspected["total_urls"],
            "completed": len(results),
            "successful": len(results) - failed,
            "failed": failed,
            "results": sorted(results, key=lambda item: item["row"]),
            "output_path": str(output_path),
        }
