"""Entrada flexible de URLs para el módulo Form Validation.

Weekly Forms conserva su matriz estricta; este adaptador independiente acepta
también un Excel sencillo con una columna URL y país opcional.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from ...form_validation.country import infer_country, normalize_allowed_url
from ..weekly_forms.spreadsheet_service import WeeklyFormsSpreadsheetService


class WeeklyLeadsSpreadsheetService(WeeklyFormsSpreadsheetService):
    """Reutiliza las reglas históricas y permite URLs QA sin columnas auxiliares."""

    # Los reportes de auditoría incluyen metadatos que no existen en la matriz
    # QA clásica. Se conservan en la fila para que el runner pueda diagnosticar
    # qué implementación de formulario encontró, sin convertirlos en campos
    # obligatorios ni cambiar el contrato histórico.
    ALIASES = {
        **WeeklyFormsSpreadsheetService.ALIASES,
        "landing_path": ("landing", "landing path", "pagina", "página"),
        "form_signal": (
            "carga formularios utel",
            "carga_formularios_utel",
            "usa backup form o html",
            "usa_backup_form_o_html",
            "formids",
            "form ids",
            "form_ids",
            "scripts forms backup",
            "scripts_forms_backup",
        ),
        "integration_hint": (
            "integracion envio",
            "integracion_envio",
            "integración envío",
            "integracion primaria",
            "integracion_primaria",
            "integración primaria",
        ),
        "source_path": ("archivo", "file", "source file"),
    }

    def preview(self, content: bytes, filename: str) -> dict[str, Any]:
        """Analiza hojas con URL obligatoria y el resto de columnas opcionales."""

        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
        sheets: list[dict[str, Any]] = []
        all_rows: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            values = list(worksheet.iter_rows(values_only=True))
            header_index = self._header_index(values)
            if header_index is None:
                continue
            headers = [self._text(value) for value in values[header_index]]
            detected = self._mapping(headers)
            if "utel_url" not in detected:
                continue
            # Conservamos los candidatos completos para que un reporte con una
            # hoja de resultados y otra hoja espejo no ejecute la misma URL dos
            # veces. La vista de cada hoja sigue mostrando solo sus pendientes.
            candidates, invalid_rows = self._flexible_rows(
                worksheet,
                header_index,
                detected,
                skip_existing=False,
            )
            all_rows.extend(candidates)
            rows = [row for row in candidates if not self._has_existing_lead(row)]
            sheets.append(
                {
                    "name": worksheet.title,
                    "headers": headers,
                    "mapping": {key: headers[index] for key, index in detected.items()},
                    "rows": rows[:200],
                    "total_rows": len(rows),
                    "invalid_rows": invalid_rows[:50],
                }
            )
        pending_unique = self._deduplicate_rows(all_rows)
        return {
            "filename": filename,
            "sheets": sheets,
            "total_rows": len(pending_unique),
            "total_invalid_rows": sum(len(sheet["invalid_rows"]) for sheet in sheets),
            "suggestions": [
                f"{sheet['name']}: {sheet['total_rows']} URLs pendientes; se ejecutarán 5 y se pausará 60 segundos."
                + (f" {len(sheet['invalid_rows'])} URLs no permitidas fueron omitidas." if sheet["invalid_rows"] else "")
                for sheet in sheets
            ],
        }

    def rows_for_mapping(self, content: bytes, mapping: dict[str, Any]) -> list[dict[str, Any]]:
        """Lee todas las URLs válidas y omite leads ya escritos en el Excel."""

        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
        selected = {
            key: self._normalize(value)
            for key, value in mapping.items()
            if isinstance(value, str) and value
        }
        if not selected.get("utel_url"):
            return []
        all_rows: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            values = list(worksheet.iter_rows(values_only=True))
            header_index = self._header_index(values)
            if header_index is None:
                continue
            headers = [self._text(value) for value in values[header_index]]
            indexes = {
                key: next(
                    (i for i, header in enumerate(headers) if self._normalize(header) == normalized),
                    None,
                )
                for key, normalized in selected.items()
            }
            if indexes.get("utel_url") is None:
                continue
            # Conservamos temporalmente las filas con Lead existente para
            # reconocerlas aunque otra hoja duplicada no tenga esa columna
            # (por ejemplo CLEAN y Hoja 1 del reporte de landings).
            all_rows.extend(
                self._flexible_rows(
                    worksheet,
                    header_index,
                    indexes,
                    skip_existing=False,
                )[0]
            )
        return self._deduplicate_rows(all_rows)

    @classmethod
    def _deduplicate_rows(cls, all_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """Elimina duplicados de hojas espejo y conserva el estado Lead existente.

        Si una copia de la misma fila ya tiene URL de lead, todas sus réplicas
        se consideran completadas. Las filas con la misma URL pero diferente
        nivel, formulario, cliente o metadatos de origen permanecen separadas.
        """

        existing_keys = {
            cls._row_identity(row)
            for row in all_rows
            if cls._has_existing_lead(row)
        }
        seen_keys: set[tuple[str, ...]] = set()
        rows: list[dict[str, Any]] = []
        for row in all_rows:
            key = cls._row_identity(row)
            if cls._has_existing_lead(row) or key in existing_keys or key in seen_keys:
                continue
            seen_keys.add(key)
            # El marcador solo sirve durante la lectura y no forma parte del
            # contrato que recibe el runner ni de los reportes descargables.
            row.pop("existing_lead_url", None)
            rows.append(row)
        return rows

    def _flexible_rows(
        self,
        worksheet: Any,
        header_index: int,
        indexes: dict[str, int | None],
        *,
        skip_existing: bool = True,
    ) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        """Construye filas normalizadas y evidencia las URLs rechazadas."""

        values = list(worksheet.iter_rows(values_only=True))
        rows: list[dict[str, Any]] = []
        invalid_rows: list[dict[str, Any]] = []
        last_country = ""
        for row_number, row_values in enumerate(values[header_index + 1 :], header_index + 2):
            raw_url = self._cell(row_values, indexes.get("utel_url"))
            if not raw_url:
                continue
            try:
                url = normalize_allowed_url(raw_url)
            except ValueError as error:
                invalid_rows.append({"row_number": row_number, "url": raw_url, "error": str(error)})
                continue

            supplied_country = self._cell(row_values, indexes.get("country"))
            if supplied_country:
                last_country = supplied_country
            elif last_country:
                supplied_country = last_country
            raw_level = self._cell(row_values, indexes.get("level"))
            level = self.infer_level(raw_level, url)
            country = infer_country(supplied_country, url, raw_level)

            existing_lead = self._cell(row_values, indexes.get("lead_url"))
            if skip_existing and existing_lead.lower().startswith(("http://", "https://")):
                continue

            raw_form = self._cell(row_values, indexes.get("form_type"))
            form_type, weekly_form_type = self._weekly_form_type(raw_form, url)
            form_signal = self._cell(row_values, indexes.get("form_signal"))
            integration_hint = self._cell(row_values, indexes.get("integration_hint"))
            rows.append(
                {
                    "sheet": worksheet.title,
                    "row_number": row_number,
                    "country": country,
                    "level": level,
                    "modality": self._cell(row_values, indexes.get("modality")) or "En linea",
                    "utel_url": url,
                    "form_type": form_type,
                    "weekly_form_type": weekly_form_type,
                    "program_name": self._cell(row_values, indexes.get("program_name")),
                    "client": self._cell(row_values, indexes.get("client")),
                    "landing_path": self._cell(row_values, indexes.get("landing_path")),
                    "form_signal": form_signal,
                    "integration_hint": integration_hint,
                    "source_path": self._cell(row_values, indexes.get("source_path")),
                    "existing_lead_url": existing_lead,
                    "inconcert_url": self._cell(row_values, indexes.get("inconcert_url")),
                    "lead_origin_url": self._cell(row_values, indexes.get("lead_origin_url")),
                    "workflow_mode": "form_validation",
                    "test_case": (
                        f"Fila {row_number} · {level} · "
                        f"{raw_form or 'Detección automática'}"
                        + (f" · {integration_hint}" if integration_hint else "")
                    ),
                }
            )
        return rows, invalid_rows

    @staticmethod
    def _row_identity(row: dict[str, Any]) -> tuple[str, ...]:
        """Identifica duplicados entre hojas sin fusionar casos distintos."""

        return tuple(
            str(row.get(field) or "").strip().casefold()
            for field in (
                "utel_url",
                "country",
                "level",
                "form_type",
                "client",
                "landing_path",
                "form_signal",
                "integration_hint",
                "source_path",
            )
        )

    @staticmethod
    def _has_existing_lead(row: dict[str, Any]) -> bool:
        """Solo una URL CRM real marca una fila como ya completada.

        Los reportes suelen escribir comentarios como ``no funciona boton`` en
        la columna Lead. Esos textos documentan un fallo, pero no deben impedir
        que Form Validation vuelva a inspeccionar la landing.
        """

        value = str(row.get("existing_lead_url") or "").strip().casefold()
        return value.startswith(("http://", "https://"))
