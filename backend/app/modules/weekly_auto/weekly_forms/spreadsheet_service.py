"""Lectura de la matriz QA usada por la automatización Weekly Forms."""

from __future__ import annotations

import io
import re
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook

from ....services.bot_spreadsheet_service import BotSpreadsheetService


class WeeklyFormsSpreadsheetService(BotSpreadsheetService):
    """Normaliza URLs y formularios sin reemplazarlos por el catálogo de PDP."""

    BATCH_SIZE = 5
    BATCH_PAUSE_SECONDS = 60
    EMERGING_MARKETS_INCONCERT = "https://mas-utel-emergentes.inconcertcc.com/mas/contact/people"

    ALIASES = {
        **BotSpreadsheetService.ALIASES,
        "lead_url": (
            "lead",
            "url lead",
            "link lead",
            "enlace lead",
            "url del lead",
            "lead url",
        ),
        "client": ("cliente", "client"),
    }

    @classmethod
    def _weekly_form_type(cls, value: str, url: str) -> tuple[str, str]:
        normalized = cls._normalize(value)
        if "form lp" in normalized or normalized in {"lp", "landing", "landing page"}:
            # Calculadora de becas conserva FooterBLC aunque algunas matrices la
            # hayan etiquetado históricamente como Form Lp.
            if "utel.edu.mx" in urlparse(url).netloc and "calculadora-de-becas" in url:
                return "footer", "footer"
            return "lateral", "form_lp"
        if "footer" in normalized or "pie" in normalized:
            return "footer", "footer"
        if "targeta" in normalized:
            # Esta etiqueta histórica corresponde a una landing clásica, no a
            # TarjetaBLC (la matriz también contiene "Tarjeta" correctamente).
            return "lateral", "form_lp"
        if "tarjeta" in normalized or "card" in normalized:
            return "tarjeta", "tarjeta"
        if "lateral" in normalized or "side" in normalized:
            return "lateral", "lateral"
        return "lateral", "form_lp"

    @classmethod
    def infer_level(cls, value: str, url: str) -> str:
        if str(value or "").strip():
            return str(value).strip()
        source = cls._normalize(url)
        rules = (
            (r"doctor", "Doctorado"),
            (r"maestr|master|posgrado", "Maestría"),
            (r"diplom|educacion-continua|curso", "Diplomado"),
            (r"bachillerato|high-school", "Bachillerato"),
            (r"licenc|carrera|pregrado|bachelor", "Licenciatura"),
        )
        for pattern, level in rules:
            if re.search(pattern, source):
                return level
        # El nivel solo sirve para escoger controles académicos. En una LP sin
        # esa información elegiremos una opción real del formulario.
        return "Licenciatura"

    @classmethod
    def effective_country(cls, country: str, level: str, url: str) -> str:
        if cls._normalize(country) != "global":
            return country
        source = cls._normalize(url)
        for token, resolved in (
            ("/usa/", "USA"),
            ("/philippines", "Filipinas"),
            ("/indonesia", "Indonesia"),
            ("/colombia", "Colombia"),
            ("/ecuador", "Ecuador"),
            ("/argentina", "Argentina"),
            ("/peru", "Peru"),
        ):
            if token in source:
                return resolved
        # Los dominios institucionales sin carpeta de país son operados desde
        # México en esta matriz (p. ej. Educación Continua).
        return "Mexico"

    @classmethod
    def default_inconcert_url(cls, country: str) -> str:
        return super().default_inconcert_url(country) or cls.EMERGING_MARKETS_INCONCERT

    def preview(self, content: bytes, filename: str) -> dict[str, Any]:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
        sheets: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            values = list(worksheet.iter_rows(values_only=True))
            header_index = self._header_index(values)
            if header_index is None:
                continue
            headers = [self._text(value) for value in values[header_index]]
            detected = self._mapping(headers)
            if not all(
                key in detected
                for key in ("country", "utel_url", "form_type", "lead_url")
            ):
                continue
            rows = self._worksheet_rows(worksheet, header_index, detected)
            sheets.append(
                {
                    "name": worksheet.title,
                    "headers": headers,
                    "mapping": {key: headers[index] for key, index in detected.items()},
                    "rows": rows[:200],
                    "total_rows": len(rows),
                }
            )
        return {
            "filename": filename,
            "sheets": sheets,
            "suggestions": [
                f"{sheet['name']}: {sheet['total_rows']} formularios pendientes; se ejecutarán 5 y se pausará 60 segundos."
                for sheet in sheets
            ],
        }

    def rows_for_mapping(self, content: bytes, mapping: dict[str, Any]) -> list[dict[str, Any]]:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
        selected = {
            key: self._normalize(value)
            for key, value in mapping.items()
            if isinstance(value, str) and value
        }
        if not all(
            selected.get(key)
            for key in ("country", "utel_url", "form_type", "lead_url")
        ):
            return []
        rows: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            values = list(worksheet.iter_rows(values_only=True))
            header_index = self._header_index(values)
            if header_index is None:
                continue
            headers = [self._text(value) for value in values[header_index]]
            indexes = {
                key: next(
                    (i for i, header in enumerate(headers) if self._normalize(header) == normalized),
                    None,
                )
                for key, normalized in selected.items()
            }
            if any(
                indexes.get(key) is None
                for key in ("country", "utel_url", "form_type", "lead_url")
            ):
                continue
            rows.extend(self._worksheet_rows(worksheet, header_index, indexes))
        return rows

    def _worksheet_rows(self, worksheet: Any, header_index: int, indexes: dict[str, int]) -> list[dict[str, Any]]:
        values = list(worksheet.iter_rows(values_only=True))
        rows: list[dict[str, Any]] = []
        last_country = ""
        for row_number, row_values in enumerate(values[header_index + 1 :], header_index + 2):
            url = self._cell(row_values, indexes.get("utel_url"))
            if not url.lower().startswith(("http://", "https://")):
                continue
            country = self._cell(row_values, indexes.get("country"))
            if country:
                last_country = country
            else:
                country = last_country
            if not country:
                continue

            existing_lead = self._cell(row_values, indexes.get("lead_url"))
            if existing_lead.lower().startswith(("http://", "https://")):
                continue

            raw_location = self._cell(row_values, indexes.get("form_type"))
            form_type, weekly_form_type = self._weekly_form_type(raw_location, url)
            level = self.infer_level(self._cell(row_values, indexes.get("level")), url)
            rows.append(
                {
                    "sheet": worksheet.title,
                    "row_number": row_number,
                    "country": country,
                    "level": level,
                    "modality": self._cell(row_values, indexes.get("modality")) or "En linea",
                    "utel_url": url,
                    "form_type": form_type,
                    "weekly_form_type": weekly_form_type,
                    "program_name": self._cell(row_values, indexes.get("program_name")),
                    "client": self._cell(row_values, indexes.get("client")),
                    "inconcert_url": self._cell(row_values, indexes.get("inconcert_url")),
                    "lead_origin_url": self._cell(row_values, indexes.get("lead_origin_url")),
                    "workflow_mode": "form_validation",
                    "test_case": f"Fila {row_number} · {level} · {raw_location or 'Form Lp'}",
                }
            )
        return rows
